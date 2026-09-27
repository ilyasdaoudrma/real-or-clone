"""Build docs/Real-or-Clone.pptx (16:9, dark, Morph transitions). python docs/make_deck.py"""
from copy import deepcopy
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches, Pt

D = Path(__file__).parent
INK, TEXT, MUTED = RGBColor(0x07, 0x08, 0x0A), RGBColor(0xF3, 0xF1, 0xEA), RGBColor(0x9A, 0x97, 0x8E)
LIME, RED, MINT, AMBER, CARD = (RGBColor(0xD4, 0xFF, 0x3A), RGBColor(0xFF, 0x4D, 0x5E), RGBColor(0x2E, 0xE5, 0x9D),
                                 RGBColor(0xFF, 0xB0, 0x20), RGBColor(0x16, 0x18, 0x1D))
FONT = "Segoe UI"
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]

MORPH = ('<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
         '<mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" Requires="p159">'
         '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="slow">'
         '<p159:morph option="byObject"/></p:transition></mc:Choice><mc:Fallback>'
         '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="slow"><p:fade/>'
         '</p:transition></mc:Fallback></mc:AlternateContent>')


def box(slide, x, y, w, h, fill=CARD, name=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.fill.background()
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    if name:
        s.name = name
    return s


def text(slide, x, y, w, h, runs, size=18, color=TEXT, bold=False, name=None, align=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame; tf.word_wrap = True
    lines = runs if isinstance(runs, list) else [runs]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align:
            p.alignment = align
        parts = line if isinstance(line, list) else [(line, color)]
        for t, c in parts:
            r = p.add_run(); r.text = t
            r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
            r.font.color.rgb = c
        p.space_after = Pt(6)
    return tb


def slide(title, kicker, n):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = INK
    # shared shapes (same names on every slide -> Morph animates them)
    box(s, 0.6 + (n % 3) * 0.2, 0.55, 0.9 + n * 0.15, 0.09, LIME, "!!bar", MSO_SHAPE.RECTANGLE)
    glow = box(s, 9.5 - n * 0.35, -1.5 + (n % 2) * 0.6, 5, 5, RGBColor(0x1B, 0x10, 0x14), "!!glow", MSO_SHAPE.OVAL)
    glow.shadow.inherit = False
    text(s, 0.6, 0.72, 8, 0.4, kicker.upper(), 12, LIME, True, "!!kicker")
    text(s, 0.6, 1.05, 12, 1.1, title, 36, TEXT, True, "!!title")
    text(s, 10.6, 6.95, 2.6, 0.4, [[("Real ", TEXT), ("or", RED), (" Clone?", TEXT)]], 12, bold=True, name="!!brand")
    s._element.append(etree.fromstring(MORPH))
    return s


def pic(s, path, x, y, w=None, h=None, name=None):
    p = s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
    if name:
        p.name = name
    return p


def stat(s, x, y, big, label, color, w=3.8):
    box(s, x, y, w, 2.0)
    text(s, x + 0.3, y + 0.2, w - 0.5, 1, big, 44, color, True)
    text(s, x + 0.3, y + 1.15, w - 0.5, 0.8, label, 14, MUTED)


shots, figs = D / "shots", D / "figures"

# 1 title
s = slide("", "GOMYCODE · Come Build with AI · Morocco · 27.09.2026", 0)
text(s, 0.6, 1.9, 12, 1.6, [[("Real ", TEXT), ("or", RED), (" Clone?", TEXT)]], 80, bold=True, name="!!hero")
text(s, 0.6, 3.45, 12, 0.8, "حقيقي أم مستنسخ؟ — detecting AI voice-clone scams in voice notes", 24, MUTED)
text(s, 0.6, 5.2, 12, 1.2, ["Ilyas Daoud · Oualid Karmoun · Ayoub El Mouhib",
                            [("Live: ", MUTED), ("real-or-clone-9m7lxdq0i.gobrev.dev", LIME)],
                            [("Code: ", MUTED), ("github.com/ilyasdaoudrma/real-or-clone", LIME)]], 16)

# 2 problem
s = slide("Voice cloning made phone scams cheap", "The problem", 1)
stat(s, 0.6, 2.4, "3 s", "of audio are enough to clone a voice", RED)
stat(s, 4.75, 2.4, "1 in 4", "adults faced an AI voice scam or know a victim", AMBER)
stat(s, 8.9, 2.4, "70%", "can't tell a clone from a real voice (77% of victims lost money)", TEXT)
text(s, 0.6, 4.8, 12, 1.5, ["“Mom, I had an accident. Send money now.” — a fake voice note, sent on WhatsApp.",
                            [("Users: ", LIME), ("families & older parents, mobile-money agents, bank / fintech call centres, fact-checkers.", TEXT)],
                            [("Source: McAfee survey, 7,054 adults, 7 countries.", MUTED)]], 16)

# 3 solution
s = slide("Upload or record a voice note. Get the truth in < 1 s.", "Our solution", 2)
pic(s, shots / "2_result.png", 0.6, 2.1, h=5.0, name="!!shot")
text(s, 7.0, 2.2, 5.8, 4.6, [[("✓ ", MINT), ("Verdict: likely real / likely clone / uncertain", TEXT)],
                             [("✓ ", MINT), ("Confidence + real-vs-clone probability bars", TEXT)],
                             [("✓ ", MINT), ("Timeline of the suspicious seconds (4 s windows)", TEXT)],
                             [("✓ ", MINT), ("3 safety tips in Arabic, French or English, written by an LLM from the verdict only — it can never change it", TEXT)],
                             [("✓ ", MINT), ("Before/after switch: XLS-R before fine-tuning vs ours", TEXT)],
                             [("✓ ", MINT), ("Clerk login, private history (audio + result, deletable) and dashboard", TEXT)]], 16)

# 4 how it works
s = slide("How we built the detector (all on one NVIDIA L40S, Brev)", "Pipeline", 3)
steps = [("Data", "FLEURS + VoxPopuli (real)\nMLAAD + In-the-Wild (fake)\nEN / FR, ~50k clips", MINT),
         ("Fresh clones", "1,619 Chatterbox clones\nof real speakers\n(scam-style texts)", RED),
         ("Voice-note sim", "Opus 16 kbps · 8 kHz\nnoise · reverb\nsame for real AND fake", AMBER),
         ("Fine-tune", "XLS-R 300M (Meta)\nreal/fake head · bf16\nbalanced 50/50", LIME),
         ("Serve", "FastAPI + web app\nGroq LLM tips\nClerk + SQLite", TEXT)]
for i, (h, b, c) in enumerate(steps):
    x = 0.6 + i * 2.5
    box(s, x, 2.5, 2.3, 3.2)
    text(s, x + 0.2, 2.65, 2.0, 0.5, h, 18, c, True)
    text(s, x + 0.2, 3.3, 2.0, 2.3, b.split("\n"), 13, TEXT)
text(s, 0.6, 6.0, 12, 0.8, "Split by speaker and by generator — no voice or TTS engine in the test set was seen in training.", 15, MUTED)

# 5 results
s = slide("Fresh clones cut the error on cloned voices 4×", "Results · held-out test (10,122 clips)", 4)
pic(s, figs / "eer_by_run.png", 0.6, 2.1, w=7.4, name="!!fig")
stat(s, 8.4, 2.1, "96.1%", "accuracy (run C) on voices it never heard", LIME, 4.3)
stat(s, 8.4, 4.35, "2.4% EER", "on ElevenLabs, OpenAI, Gemini… never seen in training", TEXT, 4.3)

# 6 shortcut
s = slide("Our own test caught a shortcut — and we fixed it", "Testing & reliability", 5)
pic(s, figs / "false_alarms_by_run.png", 0.6, 2.1, w=7.4, name="!!fig")
text(s, 8.4, 2.2, 4.5, 4.6, [[("Run A/B: ", AMBER), ("learned “sounds like FLEURS = real”. 96–99% of real voices from other sources were flagged; the team lead's own phone note scored 1.00 fake.", TEXT)],
                             [("Run C: ", MINT), ("+ varied real voices (VoxPopuli, In-the-Wild speakers) → false alarms ~7%, In-the-Wild EER 30.6% → 3.4%.", TEXT)],
                             [("Caveat: ", MUTED), ("In-the-Wild test is speaker-disjoint but no longer a new domain for run C.", MUTED)]], 15)

# 7 confusion
s = slide("Before vs after: confusion matrices", "Runs A → B → C", 6)
pic(s, figs / "confusion_matrices.png", 0.6, 2.1, w=12.1, name="!!fig")
text(s, 0.6, 6.3, 12, 0.6, "Accuracy 70.8% → 70.8% → 96.1% · F1 0.798 → 0.802 → 0.968 · 30 ms per 10 s of audio on the L40S", 15, MUTED)

# 8 app
s = slide("The app: EN / FR / AR, mobile-first, private history", "Product", 7)
pic(s, shots / "4_dashboard.png", 0.6, 2.1, h=4.9, name="!!shot")
pic(s, shots / "3_result_ar.png", 6.6, 2.1, h=4.9)
pic(s, shots / "5_mobile.png", 11.0, 2.1, h=4.9)

# 9 responsible AI
s = slide("Responsible AI, cost and limits", "Honest by design", 8)
text(s, 0.6, 2.1, 6.0, 4.8, [[("Privacy  ", LIME), ("Login required; each user's audio + results are private and deletable item by item.", TEXT)],
                             [("LLM scope  ", LIME), ("the LLM never hears the audio and cannot change the verdict; fixed tips if it fails.", TEXT)],
                             [("Consent  ", LIME), ("demo voice = team lead, written consent.", TEXT)],
                             [("Humility  ", LIME), ("“An AI can be wrong — always call back on a number you know.”", TEXT)]], 15)
text(s, 7.0, 2.1, 5.8, 4.8, [[("Cost  ", AMBER), ("1× L40S on Brev at $2.14/h — whole project ≈ $10 of the $100 credit.", TEXT)],
                             [("Limits  ", AMBER), ("English + French only — not validated on Arabic or Darija voices yet.", TEXT)],
                             [("Licences  ", AMBER), ("MLAAD is CC-BY-NC → non-commercial prototype.", TEXT)],
                             [("Next  ", AMBER), ("Darija data, real phone calls, more generators, on-device model.", TEXT)]], 15)

# 10 stack
s = slide("Stack & AI disclosure", "What we used", 9)
text(s, 0.6, 2.1, 12, 4.8, [[("Models  ", LIME), ("XLS-R 300M (Meta, Apache-2.0, fine-tuned by us) · Chatterbox Multilingual (MIT) · gpt-oss-120b on Groq (tips only)", TEXT)],
                            [("Data  ", LIME), ("FLEURS (CC-BY) · VoxPopuli (CC0) · MLAAD (CC-BY-NC) · In-the-Wild (CC-BY-SA)", TEXT)],
                            [("Infra  ", LIME), ("NVIDIA Brev L40S · PyTorch · Transformers · FastAPI · Clerk · SQLite · Motion", TEXT)],
                            [("AI help  ", LIME), ("Claude Code wrote most of the code; the team chose the approach, ran every GPU step, listened to the clones, checked the numbers and caught the shortcut.", TEXT)],
                            [("Try it  ", LIME), ("real-or-clone-9m7lxdq0i.gobrev.dev — sign in with Google or email", TEXT)]], 16)

# 11 thank you
from pptx.enum.text import PP_ALIGN
s = slide("", "", 10)
text(s, 0.6, 2.3, 12.1, 2.0, [[("Thank you", TEXT)]], 110, bold=True, name="!!hero", align=PP_ALIGN.CENTER)
text(s, 0.6, 4.35, 12.1, 0.6, [[("Real ", TEXT), ("or", RED), (" Clone?", TEXT), ("  ·  شكرًا  ·  Merci", MUTED)]], 22, align=PP_ALIGN.CENTER)
text(s, 0.6, 5.4, 12.1, 0.5, [[("real-or-clone-9m7lxdq0i.gobrev.dev", LIME)]], 16, align=PP_ALIGN.CENTER)

prs.save(D / "Real-or-Clone.pptx")
print("saved", D / "Real-or-Clone.pptx", len(prs.slides), "slides")
